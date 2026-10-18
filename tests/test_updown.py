"""The Ups and Downs event: members also guess how many titles climb and fall against last week's chart."""
import asyncio

import pytest
from discord import app_commands

import scoring
import storage
from conftest import RESULTS, TITLES
from scoring import chart_movement, compute_updown_score, updown_points

RULES = {pos: (2, 1) for pos in range(1, 11)}
UPDOWN = app_commands.Choice(name="ups and downs", value="updown")

# Week 2's chart against RESULTS (week 1's): ups = Frieren, Re:Zero, Vinland Saga; downs = One Piece, Dandadan,
# Jujutsu Kaisen; unchanged = Bleach, Spy x Family; new = Mob Psycho 100, 86; dropped = Chainsaw Man, Demon Slayer.
CHART2 = ["Frieren: Beyond Journey's End", "One Piece", "Bleach", "Re:Zero", "Dandadan", "Mob Psycho 100",
          "Jujutsu Kaisen", "Spy x Family", "Vinland Saga", "86"]


def run(coro):
    return asyncio.run(coro)


def ud(week_id, user_id):
    """(ups, downs) saved for a member, or None."""
    raw = storage.get_pick_extras(week_id, user_id)
    return (int(raw["ups"]), int(raw["downs"])) if "ups" in raw and "downs" in raw else None


async def score_week_one(env):
    """Setup channels, season, a standard week 1 with alice and bob picking its eventual chart, and score it."""
    await env.setup_channels()
    await env.start_season(TITLES)
    await env.weeks.event_template.callback(env.weeks, env.inter(env.admin), "Standard", None, "")
    await env.pick(env.alice, RESULTS)
    await env.pick(env.bob, RESULTS)
    await env.weeks.lock.callback(env.weeks, env.inter(env.admin))
    j = env.inter(env.admin)
    await env.weeks.end_week.callback(env.weeks, j, "")
    await env.submit(j, "\n".join(RESULTS))


async def open_updown(env, note=""):
    i = env.inter(env.admin)
    await env.weeks.event_template.callback(env.weeks, i, "Standard", UPDOWN, note)
    return i


# ---- the maths ---------------------------------------------------------------------------------------------------------

def test_a_guess_scores_5_when_exact_and_one_less_per_step_down_to_zero():
    assert [updown_points(g, 3) for g in range(0, 11)] == [2, 3, 4, 5, 4, 3, 2, 1, 0, 0, 0]
    assert updown_points(3, 3) == 5 and updown_points(2, 3) == updown_points(4, 3) == 4
    assert updown_points(0, 0) == 5 and updown_points(10, 0) == 0


def test_chart_movement_sorts_titles_into_up_down_same_new_and_dropped():
    m = chart_movement(RESULTS, CHART2)
    assert m["ups"] == ["Frieren: Beyond Journey's End", "Re:Zero", "Vinland Saga"]
    assert m["downs"] == ["One Piece", "Dandadan", "Jujutsu Kaisen"]
    assert m["same"] == ["Bleach", "Spy x Family"] and m["new"] == ["Mob Psycho 100", "86"]
    assert m["dropped"] == ["Chainsaw Man", "Demon Slayer"]
    assert (m["actual_ups"], m["actual_downs"]) == (3, 3)
    # spelling, case and empty spots don't confuse it
    m = chart_movement(["", "one piece", "BLEACH"], ["Bleach", "", "One Piece"])
    assert m["ups"] == ["Bleach"] and m["downs"] == ["One Piece"] and not m["new"] and not m["dropped"]
    assert chart_movement(RESULTS, RESULTS)["actual_ups"] == 0 == chart_movement(RESULTS, RESULTS)["actual_downs"]


def test_new_entries_and_drop_offs_only_count_when_those_switches_are_turned_on(monkeypatch):
    assert (chart_movement(RESULTS, CHART2)["actual_ups"], chart_movement(RESULTS, CHART2)["actual_downs"]) == (3, 3)
    monkeypatch.setattr(scoring, "COUNT_NEW_AS_UPS", True)
    assert chart_movement(RESULTS, CHART2)["actual_ups"] == 5
    monkeypatch.setattr(scoring, "COUNT_DROPPED_AS_DOWNS", True)
    assert chart_movement(RESULTS, CHART2)["actual_downs"] == 5
    assert "a new entry counts as an up" in scoring.updown_explainer("Week 1")


def test_compute_updown_score_rows_and_a_missing_guess():
    movement = chart_movement(RESULTS, CHART2)                       # actual: 3 ups, 3 downs
    total, rows = compute_updown_score((3, 1), movement)
    assert total == 5 + 3
    assert rows == [{"position": "Ups", "predicted": "3 (actual 3)", "points": 5, "reason": "exact"},
                    {"position": "Downs", "predicted": "1 (actual 3)", "points": 3, "reason": "off by 2"}]
    total, rows = compute_updown_score(None, movement)
    assert total == 0 and [r["reason"] for r in rows] == ["no prediction", "no prediction"]
    assert "Week 1" in scoring.updown_explainer("Week 1") and "new entries don't count" in scoring.updown_explainer("Week 1")


# ---- starting the event ----------------------------------------------------------------------------------------------------

def test_it_cannot_start_without_a_scored_week_before_it(env):
    async def scenario():
        await env.setup_channels()
        await env.start_season(TITLES)
        i = await open_updown(env)
        assert "can't be the first event" in i.all_text and storage.list_weeks() == []

        await env.weeks.event_template.callback(env.weeks, env.inter(env.admin), "Standard", None, "")
        await env.weeks.lock.callback(env.weeks, env.inter(env.admin))
        before = len(env.announce.sent)
        i = await open_updown(env)                                       # week 1 is locked but has no results
        assert "Week 1 has no results yet" in i.all_text and "/end-week week:week-1" in i.all_text
        assert [w["week_id"] for w in storage.list_weeks()] == ["week-1"] and len(env.announce.sent) == before
    run(scenario())


def test_the_announcement_explains_the_event_and_shows_last_weeks_chart(env):
    async def scenario():
        await score_week_one(env)
        i = await open_updown(env, "picks close Sunday")
        assert "Started **Week 2** (Ups and Downs)" in i.all_text
        assert storage.get_week("week-2")["type"] == "updown"
        content, embed, _ = env.announce.sent[-1]
        assert content == "Hello everyone! **Week 2** predictions are open!\npicks close Sunday"
        text = embed.description
        assert text.startswith("# Week 2 predictions are open") and "Type :: Ups and Downs" in text
        assert "`ups` and `downs`" in text and "5 pts" in text and "off by 1 is 4" in text and "compared with Week 1" in text
        assert "### Week 1's chart" in text and "1 :: One Piece" in text and "10 :: Vinland Saga" in text
        assert len(embed) < 6000

        j = env.inter(env.alice)
        await env.help.help_cmd.callback(env.help, j)
        assert "Ups and Downs" in j.all_text
    run(scenario())


# ---- /pick in an Ups and Downs week ------------------------------------------------------------------------------------------

def test_pick_needs_both_ups_and_downs_the_first_time_and_remembers_them_after(env):
    async def scenario():
        await score_week_one(env)
        await open_updown(env)
        uid = str(env.alice.id)

        out = await env.pick(env.alice, CHART2)
        assert "rejected" in out.all_text and "`ups` and `downs` are both needed" in out.all_text
        out = await env.pick(env.alice, CHART2, ups=3)
        assert "rejected" in out.all_text and storage.get_prediction("week-2", uid) is None
        out = await env.pick(env.alice, CHART2, ups=6, downs=5)
        assert "can't add up to more than 10" in out.all_text and storage.get_prediction("week-2", uid) is None

        posts_before = len(env.picks.sent)
        out = await env.pick(env.alice, CHART2, ups=3, downs=4)
        assert "Saved" in out.all_text and "Your Ups and Downs picks: Ups 3, Downs 4." in out.all_text
        assert ud("week-2", uid) == (3, 4) and storage.prediction_ranks(storage.get_prediction("week-2", uid)) == CHART2
        post = env.picks.sent[-1][1].description
        assert len(env.picks.sent) == posts_before + 1
        assert "### Ups and Downs" in post and "Ups :: 3" in post and "Downs :: 4" in post and "1 :: Frieren" in post

        # change a rank only: ups and downs stay. Change just ups: the ranks stay and downs stay.
        i = env.inter(env.alice)
        await env.user_cog.pick.callback(env.user_cog, i, "Bleach")
        assert ud("week-2", uid) == (3, 4) and "Saved" in i.all_text
        i = env.inter(env.alice)
        await env.user_cog.pick.callback(env.user_cog, i, *([""] * 10), 2, None)
        assert ud("week-2", uid) == (2, 4) and "ups 2" in i.all_text
        assert storage.prediction_ranks(storage.get_prediction("week-2", uid))[0] == "Bleach"
        assert len(env.picks.sent) == posts_before + 1, "edited in place"
        assert "Ups :: 2" in env.picks.sent[posts_before][2].embed.description    # the live message, edited in place

        # ups and downs may both be 0, and exactly 10 in total is fine
        assert "Saved" in (await env.pick(env.alice, CHART2, ups=0, downs=0)).all_text
        assert "Saved" in (await env.pick(env.alice, CHART2, ups=4, downs=6)).all_text
        assert ud("week-2", uid) == (4, 6)
    run(scenario())


def test_a_bad_title_or_bad_ups_refuses_everything_and_a_bare_pick_shows_the_state(env):
    async def scenario():
        await score_week_one(env)
        await open_updown(env)
        uid = str(env.alice.id)
        await env.pick(env.alice, CHART2, ups=3, downs=3)
        out = await env.pick(env.alice, ["Blech"] + CHART2[1:], ups=1, downs=1)
        assert "rejected" in out.all_text and "Rank #1 ('Blech')" in out.all_text
        assert ud("week-2", uid) == (3, 3), "the earlier ups and downs must be untouched"

        i = env.inter(env.alice)
        await env.user_cog.pick.callback(env.user_cog, i)
        guide = i.sent[-1].embed.description
        assert "Your Ups and Downs picks" in guide and "Ups :: 3" in guide and "Downs :: 3" in guide and "`ups` and `downs`" in guide
        i = env.inter(env.bob)                                            # nothing submitted yet
        await env.user_cog.pick.callback(env.user_cog, i)
        assert "Ups :: -" in i.sent[-1].embed.description and storage.get_prediction("week-2", str(env.bob.id)) is None
    run(scenario())


def test_ordinary_weeks_are_untouched_and_refuse_ups_and_downs(env):
    async def scenario():
        await env.setup_channels()
        await env.start_season(TITLES)
        await env.weeks.event_template.callback(env.weeks, env.inter(env.admin), "Standard", None, "")
        out = await env.pick(env.alice, RESULTS, ups=3, downs=3)
        assert "`ups` and `downs` don't apply in a standard week" in out.all_text and storage.get_prediction("week-1", str(env.alice.id)) is None
        out = await env.pick(env.alice, RESULTS)
        assert "Saved" in out.all_text and ud("week-1", str(env.alice.id)) is None
        assert "Ups and Downs" not in env.picks.sent[-1][1].description
        i = env.inter(env.alice)
        await env.user_cog.pick.callback(env.user_cog, i)
        assert "Your Ups and Downs picks" not in i.sent[-1].embed.description
    run(scenario())


def test_the_pick_command_offers_optional_ups_and_downs_from_0_to_10(env):
    async def scenario():
        from bot import PredictionBot
        from discord.ext import commands
        bot = PredictionBot()
        for ext in ["cogs.season", "cogs.weeks", "cogs.user", "cogs.help", "cogs.backups"]:
            await bot.load_extension(ext)
        payload = bot.tree.get_command("pick").to_dict(bot.tree)
        options = {o["name"]: o for o in payload["options"]}
        for name in ("ups", "downs"):
            o = options[name]
            assert o["type"] == 4 and not o.get("required") and (o["min_value"], o["max_value"]) == (0, 10)
            assert len(o["description"]) <= 100 and "Ups and Downs weeks only" in o["description"]
        assert [o["name"] for o in payload["options"]][-2:] == ["ups", "downs"]
        assert len(payload["options"]) == 12
        week_type = next(o for o in bot.tree.get_command("event-template").to_dict(bot.tree)["options"] if o["name"] == "week_type")
        assert "updown" in [c["value"] for c in week_type["choices"]]
        await commands.Bot.close(bot)
    run(scenario())


# ---- scoring the week -----------------------------------------------------------------------------------------------------------

def test_end_to_end_ranked_points_plus_ups_and_downs_points(env):
    async def scenario():
        await score_week_one(env)
        await open_updown(env)
        await env.pick(env.alice, CHART2, ups=3, downs=3)              # exact on both: +5 +5
        await env.pick(env.bob, RESULTS, ups=5, downs=0)               # off by 2 (+3) and by 3 (+2)
        await env.pick(env.carol, CHART2[::-1], ups=0, downs=10)       # off by 3 (+2) and by 7 (+0)
        await env.weeks.lock.callback(env.weeks, env.inter(env.admin))

        j = env.inter(env.admin)
        await env.weeks.end_week.callback(env.weeks, j, "")
        out = await env.submit(j, "\n".join(CHART2))
        assert "Results saved" in out.all_text and "scored for Week 2" in out.all_text

        scores = {s["username"]: s for s in storage.list_scores("week-2")}
        ranked = lambda picks: scoring.compute_score(picks, CHART2, RULES)[0]
        assert ranked(CHART2) == 20
        assert storage.score_points(scores["alice"]) == 20 + 10
        assert storage.score_points(scores["bob"]) == ranked(RESULTS) + 5
        assert storage.score_points(scores["carol"]) == ranked(CHART2[::-1]) + 2

        import json
        rows = json.loads(scores["bob"]["breakdown_json"])
        assert rows[-2:] == [{"position": "Ups", "predicted": "5 (actual 3)", "points": 3, "reason": "off by 2"},
                             {"position": "Downs", "predicted": "0 (actual 3)", "points": 2, "reason": "off by 3"}]

        results = next(e.description for _, e, _ in reversed(env.announce.sent) if e and e.description.startswith("# Results are in"))
        assert "Results are in: Week 2 (Ups and Downs)" in results and "### Ups and Downs vs Week 1" in results
        assert "Ups :: **3** (Frieren: Beyond Journey's End, Re:Zero, Vinland Saga)" in results
        assert "Downs :: **3** (One Piece, Dandadan, Jujutsu Kaisen)" in results
        assert "Unchanged :: Bleach, Spy x Family" in results and "New :: Mob Psycho 100, 86" in results
        assert "Dropped off :: Chainsaw Man, Demon Slayer" in results

        k = env.inter(env.bob)
        await env.user_cog.my_score.callback(env.user_cog, k, "")
        assert "Ups :: 5 (actual 3)" in k.all_text and "(off by 2)" in k.all_text and "Downs :: 0 (actual 3)" in k.all_text

        # correcting the chart re-scores with the right comparison
        fix = env.inter(env.admin)
        swapped = list(CHART2); swapped[0], swapped[1] = swapped[1], swapped[0]
        await env.weeks.end_week.callback(env.weeks, fix, "", *swapped[:2])   # frieren/one piece swapped back to week 1's order
        movement = chart_movement(RESULTS, swapped)
        assert movement["actual_ups"] == 2 and movement["actual_downs"] == 2
        # with the corrected chart: 2 ups, 2 downs. Alice guessed 3 and 3 (off by 1 each = 4 + 4).
        assert storage.score_points({s["username"]: s for s in storage.list_scores("week-2")}["alice"]) == ranked_after(swapped) + 4 + 4
    run(scenario())


def ranked_after(chart):
    return scoring.compute_score(CHART2, chart, RULES)[0]


def test_scoring_waits_for_last_weeks_results_and_never_half_saves(env):
    async def scenario():
        await score_week_one(env)
        await open_updown(env)
        await env.pick(env.alice, CHART2, ups=3, downs=3)
        await env.weeks.lock.callback(env.weeks, env.inter(env.admin))
        with storage._lock:                                              # last week's chart goes missing
            storage._db().execute("DELETE FROM results WHERE week_id = 'week-1'")

        i = env.inter(env.admin)
        await env.weeks.end_week.callback(env.weeks, i, "")
        assert i.modal is None and "Week 1 has no results yet" in i.all_text and "/end-week week:week-1" in i.all_text
        i = env.inter(env.admin)
        await env.weeks.end_week.callback(env.weeks, i, "", *CHART2[:3])
        assert "Week 1 has no results yet" in i.all_text
        assert storage.get_results("week-2") is None and storage.list_scores("week-2") == []
    run(scenario())


def test_last_weeks_chart_vanishing_while_the_form_is_open_still_saves_nothing(env):
    async def scenario():
        await score_week_one(env)
        await open_updown(env)
        await env.pick(env.alice, CHART2, ups=3, downs=3)
        await env.weeks.lock.callback(env.weeks, env.inter(env.admin))

        j = env.inter(env.admin)
        await env.weeks.end_week.callback(env.weeks, j, "")             # the form opens: last week's chart is there
        assert j.modal is not None
        with storage._lock:                                              # ...then it goes missing before the host submits
            storage._db().execute("DELETE FROM results WHERE week_id = 'week-1'")
        out = await env.submit(j, "\n".join(CHART2))
        assert "Week 1 has no results yet" in out.all_text and "Nothing was saved" in out.all_text
        assert storage.get_results("week-2") is None and storage.list_scores("week-2") == []
    run(scenario())


def test_reset_clears_ups_and_downs_and_the_export_includes_them(env):
    async def scenario():
        import io
        import openpyxl
        await score_week_one(env)
        await open_updown(env)
        await env.pick(env.alice, CHART2, ups=3, downs=3)
        i = env.inter(env.admin)
        await env.backups.export.callback(env.backups, i)
        wb = openpyxl.load_workbook(io.BytesIO(i.sent[-1].file.fp.read()))
        assert "pick_extras" in wb.sheetnames
        assert [tuple(r) for r in wb["pick_extras"].iter_rows(values_only=True)][1:] == [
            ("week-2", str(env.alice.id), "ups", "3"), ("week-2", str(env.alice.id), "downs", "3")]

        storage.reset_season()
        assert ud("week-2", str(env.alice.id)) is None and storage.list_pick_extras("week-2") == {}
    run(scenario())
