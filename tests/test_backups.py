import asyncio
import io
import os
import re
import sqlite3

import openpyxl
import pytest

import backup
import storage
from conftest import BACKUP_CHANNEL, RESULTS, TITLES
from fakes import FakeAttachment


def run(coro):
    return asyncio.run(coro)


def backup_messages(env):
    return sorted((m for m in env.backups_channel.messages.values() if m.attachments), key=lambda m: m.id)


def kept(env):
    return [m for m in backup_messages(env) if m.attachments[0].filename.endswith("-keep.db")]


def age_last_backup(days):
    """Pretends the last backup was posted `days` days ago."""
    from datetime import datetime, timedelta, timezone
    stamp = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat(timespec="seconds")
    storage._db().execute("INSERT OR REPLACE INTO meta(key, value) VALUES ('last_backup_at', ?)", (stamp,))


def zeros():
    return {"weeks": 0, "submissions": 0, "titles": 0, "scored_rows": 0}


async def restore_via_command(env, file=None, click=0):
    """Runs /restore, presses a button (0 = confirm, 1 = cancel) and returns the interaction."""
    task, i = await env.start_command(lambda i: env.backups.restore.callback(env.backups, i, file))
    if not task.done():
        await env.click(env.confirm_view(i), click)
    await task
    return i


# ---- automatic backups -----------------------------------------------------------------------------------

def test_scoring_a_week_posts_a_backup_and_the_schedule_is_weekly(env):
    async def scenario():
        await env.play_week()                                    # /end-week posts a backup of the scored week
        messages = backup_messages(env)
        assert len(messages) == 1 and "week-1 scored" in messages[0].content
        attachment = messages[0].attachments[0]
        assert re.fullmatch(r"prediction-backup-\d{8}-\d{6}\.db", attachment.filename)
        info = storage.inspect_snapshot(await attachment.read())
        assert (info["weeks"], info["submissions"], info["titles"]) == (1, 2, 13)
        assert not storage.is_dirty()

        storage.add_anime("Late Addition")                       # changed, but the last backup is fresh
        await env.backups.save_if_due()
        assert len(backup_messages(env)) == 1 and storage.is_dirty()

        age_last_backup(6)
        await env.backups.save_if_due()
        assert len(backup_messages(env)) == 1, "6 days is not a week yet"

        age_last_backup(8)
        await env.backups.save_if_due()
        assert len(backup_messages(env)) == 2 and "scheduled" in backup_messages(env)[-1].content
        assert not storage.is_dirty()

        age_last_backup(30)                                      # old, but nothing changed: nothing to save
        await env.backups.save_if_due()
        assert len(backup_messages(env)) == 2
    run(scenario())


def test_the_first_ever_change_is_backed_up_without_waiting_a_week(env):
    async def scenario():
        storage.add_anime("First")
        assert storage.last_backup_at() is None
        await env.backups.save_if_due()
        assert len(backup_messages(env)) == 1
    run(scenario())


def test_old_automatic_backups_are_pruned_but_milestones_never_are(env):
    async def scenario():
        for n in range(8):
            storage.add_anime(f"Show {n}")
            age_last_backup(8)
            await env.backups.save_if_due()
        assert len(backup_messages(env)) == 5, "BACKUP_KEEP is 5 in the tests"

        i = env.inter(env.admin)
        await env.backups.backup_now.callback(env.backups, i)
        assert "Backup posted" in i.all_text
        for n in range(8, 16):
            storage.add_anime(f"Show {n}")
            age_last_backup(8)
            await env.backups.save_if_due()

        names = [m.attachments[0].filename for m in backup_messages(env)]
        assert len([n for n in names if n.endswith("-keep.db")]) == 1, "the manual backup must survive pruning"
        assert len([n for n in names if not n.endswith("-keep.db")]) == 5
    run(scenario())


def test_a_failing_backup_channel_is_reported_not_fatal_and_recovers(env):
    async def scenario():
        await env.setup_channels()
        env.backups_channel.forbid_send = True
        storage.add_anime("Something")
        await env.backups.save_if_due()                       # must not raise
        assert "Attach Files" in backup.state.last_error and storage.is_dirty()

        i = env.inter(env.admin)
        await env.help.help_admin.callback(env.help, i)
        assert "[todo] Backups" in i.all_text and "last attempt failed" in i.all_text
        i = env.inter(env.admin)
        await env.backups.backup_now.callback(env.backups, i)
        assert "The backup failed" in i.all_text

        env.backups_channel.forbid_send = False
        await env.backups.save_if_due()
        assert backup.state.last_error is None and not storage.is_dirty()
        i = env.inter(env.admin)
        await env.help.help_admin.callback(env.help, i)
        assert "[ok] Backups" in i.all_text and "up to date" in i.all_text
        assert len(i.last.embed) < 6000
    run(scenario())


def test_backup_needs_the_right_channel_permissions_to_be_called_ok(env):
    async def scenario():
        from fakes import FakePerms
        env.backups_channel._perms = FakePerms(attach_files=False, read_message_history=False)
        i = env.inter(env.admin)
        await env.help.help_admin.callback(env.help, i)
        assert "attach files" in i.all_text and "read message history" in i.all_text
    run(scenario())


# ---- /restore ------------------------------------------------------------------------------------------------------

def test_restore_latest_brings_the_season_back_after_saving_the_current_state(env):
    async def scenario():
        await env.play_week()
        await env.backups.save_if_due()
        good = storage.data_summary()
        storage.reset_season()                                  # the disaster
        storage.add_anime("Junk")

        task, i = await env.start_command(lambda i: env.backups.restore.callback(env.backups, i, None))
        prompt = i.original_edits[-1]
        assert "newest backup" in prompt and "1 week(s)" in prompt and "replaces everything" in prompt
        await env.click(env.confirm_view(i), 0)
        await task

        assert storage.data_summary() == good
        assert "Restored" in i.original_edits[-1] and "this backup" in i.original_edits[-1]
        safety = kept(env)
        assert len(safety) == 1 and "before restore" in safety[0].content
        assert storage.inspect_snapshot(await safety[0].attachments[0].read())["titles"] == 1, "undo copy holds the pre-restore state"
        assert storage.is_dirty(), "restored state should be re-posted as the newest backup"

        assert storage.get_config("announcement_channel_id") == str(env.announce.id)
        j = env.inter(env.alice)
        await env.user_cog.leaderboard.callback(env.user_cog, j, "")
        assert "alice" in j.last.embed.description
    run(scenario())


def test_restore_from_an_attached_file(env):
    async def scenario():
        await env.play_week()
        data, _ = storage.snapshot()
        storage.reset_season()
        i = await restore_via_command(env, FakeAttachment("my-old-backup.db", data))
        assert "your file `my-old-backup.db`" in i.original_edits[0]
        assert storage.data_summary()["weeks"] == 1 and len(storage.get_anime_list()) == 13
    run(scenario())


def test_cancelling_a_restore_changes_nothing_and_strangers_cannot_confirm(env):
    async def scenario():
        await env.play_week()
        await env.backups.save_if_due()
        storage.reset_season()
        task, i = await env.start_command(lambda i: env.backups.restore.callback(env.backups, i, None))
        view = env.confirm_view(i)
        stranger = env.inter(env.alice)
        assert await view.interaction_check(stranger) is False and "someone else" in stranger.all_text
        cancel = await env.click(view, 1)
        await task
        assert "Cancelled" in cancel.last.content
        assert storage.data_summary() == zeros() and not kept(env)
    run(scenario())


def test_restore_refuses_bad_files_without_touching_anything(env):
    async def scenario():
        storage.add_anime("Precious")
        other = os.path.join(os.environ["DATA_DIR"], "other.db")
        conn = sqlite3.connect(other)
        conn.execute("CREATE TABLE notes(x)")
        conn.commit()
        conn.close()
        for name, data in (("photo.db", b"\x89PNG not a database" * 20), ("other.db", open(other, "rb").read())):
            i = env.inter(env.admin)
            await env.backups.restore.callback(env.backups, i, FakeAttachment(name, data))
            assert "can't use that file" in i.all_text and not any(s.view for s in i.sent), name
        assert storage.get_anime_list() == ["Precious"]
    run(scenario())


def test_restore_with_no_backups_says_so(env):
    async def scenario():
        i = env.inter(env.admin)
        await env.backups.restore.callback(env.backups, i, None)
        assert "no usable backups with data" in i.all_text
        env.backups_channel.forbid_history = True
        i = env.inter(env.admin)
        await env.backups.restore.callback(env.backups, i, None)
        assert "Read Message History" in i.all_text
    run(scenario())


def test_restore_is_aborted_if_the_safety_backup_cannot_be_saved(env):
    async def scenario():
        await env.play_week()
        await env.backups.save_if_due()
        storage.add_anime("After the backup")
        env.backups_channel.forbid_send = True
        task, i = await env.start_command(lambda i: env.backups.restore.callback(env.backups, i, None))
        await env.click(env.confirm_view(i), 0)
        await task
        assert "nothing was changed" in i.original_edits[-1]
        assert "After the backup" in storage.get_anime_list(), "current data must be left alone"
    run(scenario())


def test_damaged_or_unrelated_files_in_the_channel_are_skipped(env):
    async def scenario():
        await env.play_week()
        await env.backups.save_if_due()
        env.backups_channel.plant("prediction-backup-29991231-235959.db", b"corrupt" * 100)
        env.backups_channel.plant("cat-picture.png", b"\x89PNG")
        message, data, info = await backup.latest_valid_backup(env.client)
        assert info["weeks"] == 1 and message.attachments[0].filename != "prediction-backup-29991231-235959.db"
    run(scenario())


# ---- fresh install: automatic restore --------------------------------------------------------------------------------------

def test_a_wiped_host_restores_itself_from_the_newest_backup(env, tmp_path):
    async def scenario():
        await env.play_week()
        await env.backups.save_if_due()
        before = storage.data_summary()
        path = str(tmp_path / "fresh-disk" / "prediction.db")           # the host lost its disk

        assert await backup.prepare_storage(env.client, path) == "restored"
        assert storage.data_summary() == before
        assert storage.get_config("announcement_channel_id") == str(env.announce.id)
        assert len(storage.list_templates()) == 2
        assert await backup.prepare_storage(env.client, path) == "existing", "a normal restart must not restore again"
    run(scenario())


def test_first_ever_start_with_no_backups_begins_empty(env, tmp_path):
    async def scenario():
        path = str(tmp_path / "prediction.db")
        assert await backup.prepare_storage(env.client, path) == "new"
        assert storage.data_summary() == zeros() and len(storage.list_templates()) == 2
        env.backups_channel.plant("prediction-backup-20260101-000000.db", b"garbage" * 50)
        path2 = str(tmp_path / "second" / "prediction.db")
        assert await backup.prepare_storage(env.client, path2) == "new", "only damaged backups: start empty, don't crash"
    run(scenario())


@pytest.mark.parametrize("break_channel", ["missing", "history"])
def test_fresh_start_refuses_to_run_blind_and_leaves_no_empty_database(env, tmp_path, break_channel):
    async def scenario():
        await env.play_week()
        await env.backups.save_if_due()
        if break_channel == "missing":
            env.client.channels.pop(BACKUP_CHANNEL)
        else:
            env.backups_channel.forbid_history = True
        path = str(tmp_path / "prediction.db")
        with pytest.raises(backup.BackupError, match="fresh install"):
            await backup.prepare_storage(env.client, path)
        assert not os.path.exists(path), "an empty file would make the next start skip the restore"
    run(scenario())


# ---- /reset ---------------------------------------------------------------------------------------------------------------------

def test_reset_posts_a_protected_backup_first_and_restore_undoes_it(env):
    async def scenario():
        await env.play_week()
        good = storage.data_summary()
        posted_before = len(backup_messages(env))

        task, i = await env.start_command(lambda i: env.season.reset.callback(env.season, i))
        assert "backup file" in i.sent[0].text and "1 week(s)" in i.sent[0].text
        await env.click(env.confirm_view(i), 1)
        await task
        assert storage.data_summary() == good and len(backup_messages(env)) == posted_before, "cancel must change nothing"

        task, i = await env.start_command(lambda i: env.season.reset.callback(env.season, i))
        await env.click(env.confirm_view(i), 0)
        await task
        assert storage.data_summary() == zeros() and "[this backup]" in i.original_edits[-1]
        assert storage.get_config("announcement_channel_id") == str(env.announce.id)
        assert len(kept(env)) == 1 and "before reset" in kept(env)[0].content

        age_last_backup(8)
        await env.backups.save_if_due()                      # a week later the schedule posts the now-empty state
        newest = backup_messages(env)[-1]
        assert storage.inspect_snapshot(await newest.attachments[0].read())["weeks"] == 0
        i = await restore_via_command(env, None)
        assert "newest backup that has data" in i.original_edits[0]
        assert storage.data_summary()["weeks"] == 1, "restore must skip the empty post-reset backup"
    run(scenario())


def test_reset_refuses_to_clear_anything_if_the_backup_cannot_be_saved(env):
    async def scenario():
        await env.play_week()
        good = storage.data_summary()
        env.backups_channel.forbid_send = True
        task, i = await env.start_command(lambda i: env.season.reset.callback(env.season, i))
        await env.click(env.confirm_view(i), 0)
        await task
        assert "nothing was cleared" in i.original_edits[-1]
        assert storage.data_summary() == good
    run(scenario())


# ---- /export -------------------------------------------------------------------------------------------------------------------

def test_export_is_a_private_excel_file_with_standings_and_text_only_cells(env):
    async def scenario():
        await env.play_week()
        storage.set_ballot_picks("week-1", "9", "zed", {"Anime": "=1+1"})
        i = env.inter(env.admin)
        await env.backups.export.callback(env.backups, i)
        sent = i.sent[-1]
        assert sent.ephemeral and sent.file.filename.endswith(".xlsx")

        wb = openpyxl.load_workbook(io.BytesIO(sent.file.fp.read()))
        assert wb.sheetnames[0] == "standings" and {"weeks", "predictions", "scores", "ballot_picks"} <= set(wb.sheetnames)
        rows = list(wb["standings"].iter_rows(values_only=True))
        assert rows[0] == ("rank", "username", "season_points")
        assert {r[1]: r[2] for r in rows[1:]} == {"alice": 20, "bob": 20}
        cell = next(c for row in wb["ballot_picks"].iter_rows() for c in row if c.value == "=1+1")
        assert cell.data_type == "s", "a member's text must never become an Excel formula"
    run(scenario())


# ---- /set-points ----------------------------------------------------------------------------------------------------------------

def test_set_points_changes_scoring_and_asks_for_a_rescore(env):
    async def scenario():
        await env.setup_channels()
        await env.start_season(TITLES)
        await env.weeks.event_template.callback(env.weeks, env.inter(env.admin), "Standard", None, "")
        i = env.inter(env.admin)
        await env.weeks.set_points.callback(env.weeks, i, 5, 2, "")
        assert "5 pts exact / 2 pts wrong spot" in i.all_text and "re-score" not in i.all_text
        assert storage.get_rules("week-1") == {p: (5, 2) for p in range(1, 11)}

        j = env.inter(env.alice)
        await env.user_cog.pick.callback(env.user_cog, j)
        await env.submit(j, "\n".join(RESULTS))
        await env.weeks.lock.callback(env.weeks, env.inter(env.admin))
        k = env.inter(env.admin)
        await env.weeks.end_week.callback(env.weeks, k, "")
        await env.submit(k, "\n".join(RESULTS))
        assert storage.score_points(storage.list_scores("week-1")[0]) == 50

        i = env.inter(env.admin)
        await env.weeks.set_points.callback(env.weeks, i, 3, 1, "week-1")
        assert "run `/end-week` to re-score" in i.all_text
        k = env.inter(env.admin)
        await env.weeks.end_week.callback(env.weeks, k, "")
        await env.submit(k, "\n".join(RESULTS))
        assert storage.score_points(storage.list_scores("week-1")[0]) == 30

    run(scenario())
