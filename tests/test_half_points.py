"""Half points (from a 1.5x multiplier) survive storage, totals, display, backups and the Excel export."""
import asyncio
import io

import openpyxl

import storage
from common import season_totals
from style import pts


def run(coro):
    return asyncio.run(coro)


def test_pts_prints_half_points_and_never_shows_5_point_0():
    assert [pts(n) for n in (1, 1.0, -1, 0, 2, 5.0, 1.5, 7.5, -0.5)] == [
        "1 pt", "1 pt", "-1 pt", "0 pts", "2 pts", "5 pts", "1.5 pts", "7.5 pts", "-0.5 pts"]


def test_half_points_are_stored_summed_and_ranked_exactly(env):
    async def scenario():
        await env.setup_channels()
        await env.start_season(["a1", "b2", "c3", "d4", "e5", "f6", "g7", "h8", "i9", "j10"])
        storage.replace_scores("week-1", [
            {"user_id": "1", "username": "alice", "points": 7.5, "breakdown": []},
            {"user_id": "2", "username": "bob", "points": 7, "breakdown": []},
        ])
        storage.replace_scores("week-2", [
            {"user_id": "1", "username": "alice", "points": 1.5, "breakdown": []},
            {"user_id": "2", "username": "bob", "points": 2.0, "breakdown": []},
        ])
        scores = storage.list_scores()
        assert [storage.score_points(s) for s in scores] == [7.5, 7, 1.5, 2]
        assert all(isinstance(storage.score_points(s), int) for s in scores if s["username"] == "bob")
        totals, _ = season_totals(scores)
        assert totals == {"1": 9, "2": 9} and all(isinstance(v, int) for v in totals.values()), "9.0 must show as 9"

        storage.replace_scores("week-3", [{"user_id": "1", "username": "alice", "points": 0.5, "breakdown": []}])
        totals, _ = season_totals(storage.list_scores())
        assert totals == {"1": 9.5, "2": 9}

        i = env.inter(env.alice)
        await env.user_cog.leaderboard.callback(env.user_cog, i, "")
        assert "**alice** — 9.5 pts" in i.all_text and "**bob** — 9 pts" in i.all_text
    run(scenario())


def test_half_points_survive_the_backup_file_and_the_excel_export(env):
    async def scenario():
        await env.setup_channels()
        await env.start_season(["a1", "b2", "c3", "d4", "e5", "f6", "g7", "h8", "i9", "j10"])
        await env.weeks.event_template.callback(env.weeks, env.inter(env.admin), "Standard", None, "")
        storage.replace_scores("week-1", [
            {"user_id": "1", "username": "alice", "points": 7.5, "breakdown": []},
            {"user_id": "2", "username": "bob", "points": 2, "breakdown": []},
        ])
        storage.replace_scores("week-2", [{"user_id": "1", "username": "alice", "points": 1.5, "breakdown": []}])

        i = env.inter(env.admin)
        await env.backups.export.callback(env.backups, i)
        wb = openpyxl.load_workbook(io.BytesIO(i.sent[-1].file.fp.read()))
        rows = {r[1]: r[2] for r in list(wb["standings"].iter_rows(values_only=True))[1:]}
        assert rows == {"alice": 9, "bob": 2}
        week_one = {r[2]: r[3] for r in list(wb["scores"].iter_rows(values_only=True))[1:] if r[0] == "week-1"}
        assert week_one == {"alice": 7.5, "bob": 2}

        snapshot, _ = storage.snapshot()                    # what a backup file holds
        storage.replace_scores("week-1", [])
        storage.restore_snapshot(snapshot)
        assert sorted(storage.score_points(s) for s in storage.list_scores("week-1")) == [2, 7.5]
    run(scenario())
