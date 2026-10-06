"""Builds the Excel export for /export from storage.dump_tables()."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font

from textutil import num


def _standings(tables: dict) -> list:
    headers, rows = tables["scores"]
    user, name, points = headers.index("user_id"), headers.index("username"), headers.index("points")
    totals, names = {}, {}
    for row in rows:
        totals[row[user]] = totals.get(row[user], 0) + float(row[points])
        names[row[user]] = row[name]
    ordered = sorted(totals.items(), key=lambda kv: (-kv[1], names[kv[0]]))
    return [(rank, names[uid], num(pts)) for rank, (uid, pts) in enumerate(ordered, start=1)]


def _write(ws, headers: list, rows: list):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                # Anything typed by members stays text, never a formula (a guess like "=1+1").
                cell.data_type = "s"
    ws.freeze_panes = "A2"


def build_xlsx(tables: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _write(wb.create_sheet("standings"), ["rank", "username", "season_points"], _standings(tables))
    for name, (headers, rows) in tables.items():
        _write(wb.create_sheet(name[:31]), headers, rows)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
